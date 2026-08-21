from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO
from typing import Iterable
from urllib.parse import quote

import frappe
from dateutil import parser as date_parser
from google.auth.transport.requests import AuthorizedSession
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from tap_lms.onboarding.backend_upload_utils import (
    FAILED_ROWS_GCP_PROJECT_ID,
    GCP_CREDENTIALS_PROJECT_ID,
    GLIFIC_CSV_HEADERS,
    GOOGLE_SHEETS_READWRITE_SCOPES,
    get_google_service_account_credentials as _get_google_service_account_credentials,
    upload_bytes_to_gcs as _upload_bytes_to_gcs,
    upload_glific_contact_csv as _upload_glific_contact_csv,
)
from tap_lms.tap_lms.doctype.student.student import _reserve_next_student_name


FIXED_STUDENT_REGISTRATION_SHEETS = [
    {
        "language": "English",
        "spreadsheet_id": "1bePUEi1KaXJfRB4O4KggkaIxDH9l7OpMZPBD9pBTgcg",
        "sheet_id": 0,
    },
    {
        "language": "Hinglish",
        "spreadsheet_id": "1GY3-zhGzAFHbjHy65FyrLh1EyAdoNNifsJPB4RwE9YE",
        "sheet_id": 0,
    },
    {
        "language": "Hindi",
        "spreadsheet_id": "1sJ-FsyNcI5yeyXeAVhFpNGJepNPWJ2gBTC6BYOkilr8",
        "sheet_id": 0,
    },
    {
        "language": "Punjabi",
        "spreadsheet_id": "13bx8lXNT_0yi4cKbywOiVxiyUufoUujZmpbF4adSUCs",
        "sheet_id": 0,
    },
    {
        "language": "Marathi",
        "spreadsheet_id": "1QH7Hi7tqe8c-ABAD5spIfA4dXJjxuOhzukua5nQP98c",
        "sheet_id": 0,
    },
]

SOURCE_COLUMNS = {
    "timestamp": "timestamp",
    "contact_phone_number": "contact_phone_number",
    "gender": "gender",
    "grade": "grade",
    "shall_we_begin": "shall_we_begin",
    "student_name": "student_name",
}
REGISTRATION_STATUS_COLUMN = "registration_status"
PROCESS_STATUS_COLUMN = "process status"

STATUS_DONE = "Done"
STATUS_PREPARED = "Prepared"
STATUS_NOT_AGREE = "Skipped: shall_we_begin does not contain Agree"
DUPLICATE_DONE_MESSAGE = "Duplicate contact_phone_number already registered; student not imported"
DUPLICATE_RUN_MESSAGE = "Duplicate contact_phone_number in current run; first valid row will be imported"
PROCESS_STATUS_COMPLETE = "complete"
PROCESS_STATUS_FAIL = "fail"
GLIFIC_CONTACT_FIELD_SCHOOL_ID = "school_id"

PREPARED_HEADERS = [
    "Language",
    "Spreadsheet Title",
    "Sheet Title",
    "Source Row",
    "Timestamp",
    "Student Name",
    "Contact Phone Number",
    "Canonical Phone",
    "Gender",
    "Grade",
    "School ID",
    "Batch",
    "Course Vertical",
    "Course Names",
    "Level",
    "Prepare Status",
    "Message",
]
NOT_DONE_ROWS_CSV_HEADERS = [
    *PREPARED_HEADERS,
    "Registration Status",
    "Process Status",
]


@dataclass(frozen=True)
class SourceSheet:
    language: str
    spreadsheet_id: str
    sheet_id: int
    spreadsheet_title: str
    sheet_title: str
    header_map: dict[str, int]
    status_column_index: int
    process_status_column_index: int
    rows: list[dict]


def prepare_student_sheet_registration(log_fn=None, progress_fn=None) -> dict:
    session = _get_sheets_session()
    sheets = _read_all_source_sheets(session, ensure_status_column=True)

    done_phones = _collect_done_phones(sheets)
    prepared_rows: list[dict] = []
    status_updates: list[dict] = []
    ready_phones: set[str] = set()
    school_lookup = _GlificSchoolLookup()

    raw_count = 0
    skipped_done = 0
    for sheet in sheets:
        language_id = _get_language_id(sheet.language)
        if not language_id:
            raise frappe.ValidationError(f"TAP Language not found: {sheet.language}")

        for source_row in sheet.rows:
            if _row_is_blank(source_row):
                continue

            raw_count += 1
            current_status = source_row.get("registration_status") or ""
            if _status_is_done(current_status):
                skipped_done += 1
                status_updates.append(
                    _process_status_update_from_source(sheet, source_row, PROCESS_STATUS_COMPLETE)
                )
                continue

            prepared = _prepare_source_row(
                source_row,
                sheet,
                language_id,
                done_phones,
                ready_phones,
                school_lookup,
            )
            prepared_rows.append(prepared)

            if prepared["prepare_status"] == "Ready":
                ready_phones.add(prepared["phone"])
                status_updates.append(_status_update(prepared, STATUS_PREPARED))
                status_updates.append(_process_status_update(prepared, PROCESS_STATUS_COMPLETE))
            elif prepared.get("message"):
                status_updates.append(_status_update(prepared, prepared["message"]))
                status_updates.append(_process_status_update(prepared, _process_status_for_row(prepared)))

    if status_updates:
        _write_status_updates(session, status_updates)

    prepared_file_url = _create_and_upload_prepared_workbook(prepared_rows)
    failed_rows = [row for row in prepared_rows if row.get("prepare_status") != "Ready"]
    failed_rows_file_url = _create_and_upload_failed_rows_workbook(failed_rows) if failed_rows else ""
    not_done_rows_file_url = _create_and_upload_not_done_rows_csv(failed_rows) if failed_rows else ""

    summary = {
        "raw_rows": raw_count,
        "skipped_done_rows": skipped_done,
        "prepared_rows": len([row for row in prepared_rows if row.get("prepare_status") == "Ready"]),
        "failed_rows": len(failed_rows),
        "duplicate_rows": len([
            row for row in failed_rows
            if str(row.get("message") or "").startswith("Duplicate contact_phone_number")
        ]),
        "prepared_file_url": prepared_file_url,
        "failed_rows_file_url": failed_rows_file_url,
        "not_done_rows_file_url": not_done_rows_file_url,
        "sheets": [
            {
                "language": sheet.language,
                "spreadsheet_id": sheet.spreadsheet_id,
                "spreadsheet_title": sheet.spreadsheet_title,
                "sheet_title": sheet.sheet_title,
            }
            for sheet in sheets
        ],
    }
    _emit(log_fn, f"[student-sheet-registration] prepared summary={summary}")
    _emit_progress(progress_fn, {"event": "prepared", "summary": dict(summary)})
    return {
        "summary": summary,
        "prepared_rows": prepared_rows,
    }


def upload_prepared_student_sheet_registration(
    prepared_rows: list[dict],
    import_user: str | None = None,
    log_fn=None,
    progress_fn=None,
) -> dict:
    import_user = (import_user or frappe.session.user or "Administrator").strip()
    session = _get_sheets_session()
    current_sheets = _read_all_source_sheets(session, ensure_status_column=True)
    current_rows_by_key = _index_current_rows(current_sheets)
    done_phones = _collect_done_phones(current_sheets)

    ready_rows = [row for row in prepared_rows if row.get("prepare_status") == "Ready"]
    success_rows: list[dict] = []
    failed_rows: list[dict] = [
        dict(row) for row in prepared_rows if row.get("prepare_status") != "Ready"
    ]
    status_updates: list[dict] = []
    uploaded_phones: set[str] = set()

    for index, row in enumerate(ready_rows, start=1):
        _emit_progress(progress_fn, {
            "event": "upload_progress",
            "processed": index - 1,
            "total": len(ready_rows),
        })

        current_row = current_rows_by_key.get(_source_row_key(row))
        if current_row and current_row.get("process_status_range") and not row.get("process_status_range"):
            row["process_status_range"] = current_row["process_status_range"]

        current_status = (current_row or {}).get("registration_status") or ""
        current_phone = _canonicalize_phone((current_row or {}).get("contact_phone_number"))

        if not current_row:
            failed_rows.append(_failed_copy(row, "Source row not found; re-run Prepare Data"))
            status_updates.append(_status_update(row, "Source row not found; re-run Prepare Data"))
            status_updates.append(_process_status_update(row, PROCESS_STATUS_FAIL))
            continue
        if current_phone != row.get("phone"):
            message = "Source row changed after prepare; re-run Prepare Data"
            failed_rows.append(_failed_copy(row, message))
            status_updates.append(_status_update(row, message))
            status_updates.append(_process_status_update(row, PROCESS_STATUS_FAIL))
            continue
        if _status_is_done(current_status):
            status_updates.append(_process_status_update(row, PROCESS_STATUS_COMPLETE))
            continue
        if row["phone"] in done_phones:
            failed_rows.append(_failed_copy(row, DUPLICATE_DONE_MESSAGE))
            status_updates.append(_status_update(row, DUPLICATE_DONE_MESSAGE))
            status_updates.append(_process_status_update(row, PROCESS_STATUS_COMPLETE))
            continue
        if row["phone"] in uploaded_phones:
            failed_rows.append(_failed_copy(row, DUPLICATE_RUN_MESSAGE))
            status_updates.append(_status_update(row, DUPLICATE_RUN_MESSAGE))
            status_updates.append(_process_status_update(row, PROCESS_STATUS_COMPLETE))
            continue

        savepoint = f"ssr_{index}"
        frappe.db.savepoint(savepoint)
        try:
            student = _upsert_student(row, import_user=import_user)
            frappe.db.commit()

            success_row = dict(row)
            success_row["student_id"] = student.name
            success_rows.append(success_row)
            uploaded_phones.add(row["phone"])
            done_phones.add(row["phone"])
            status_updates.append(_status_update(row, STATUS_DONE))
            status_updates.append(_process_status_update(row, PROCESS_STATUS_COMPLETE))
        except Exception as exc:
            frappe.db.rollback(save_point=savepoint)
            message = _short_error(exc)
            failed_rows.append(_failed_copy(row, message))
            status_updates.append(_status_update(row, message))
            status_updates.append(_process_status_update(row, PROCESS_STATUS_FAIL))
            frappe.db.commit()
            _emit(log_fn, f"[student-sheet-registration] row_failed row={row.get('row_number')} error={message}")

    if status_updates:
        _write_status_updates(session, status_updates)

    failed_rows_file_url = _create_and_upload_failed_rows_workbook(failed_rows) if failed_rows else ""
    not_done_rows_file_url = _create_and_upload_not_done_rows_csv(failed_rows) if failed_rows else ""
    glific_contact_files = _create_and_upload_glific_contact_csvs(success_rows)
    summary = {
        "prepared_rows": len(ready_rows),
        "uploaded_rows": len(success_rows),
        "failed_rows": len(failed_rows),
        "duplicate_rows": len([
            row for row in failed_rows
            if str(row.get("message") or "").startswith("Duplicate contact_phone_number")
        ]),
        "failed_rows_file_url": failed_rows_file_url,
        "not_done_rows_file_url": not_done_rows_file_url,
        "glific_contact_files": glific_contact_files,
    }
    _emit(log_fn, f"[student-sheet-registration] uploaded summary={summary}")
    _emit_progress(progress_fn, {"event": "uploaded", "summary": dict(summary)})
    return summary


def _prepare_source_row(
    source_row: dict,
    sheet: SourceSheet,
    language_id: str,
    done_phones: set[str],
    ready_phones: set[str],
    school_lookup: "_GlificSchoolLookup | None" = None,
) -> dict:
    base = {
        "language": sheet.language,
        "language_id": language_id,
        "spreadsheet_id": sheet.spreadsheet_id,
        "spreadsheet_title": sheet.spreadsheet_title,
        "sheet_id": sheet.sheet_id,
        "sheet_title": sheet.sheet_title,
        "row_number": source_row["row_number"],
        "status_column_index": sheet.status_column_index,
        "status_range": _cell_range(sheet.sheet_title, sheet.status_column_index, source_row["row_number"]),
        "process_status_column_index": sheet.process_status_column_index,
        "process_status_range": _cell_range(
            sheet.sheet_title,
            sheet.process_status_column_index,
            source_row["row_number"],
        ),
        "timestamp": source_row.get("timestamp") or "",
        "contact_phone_number": source_row.get("contact_phone_number") or "",
        "student_name_raw": source_row.get("student_name") or "",
        "gender_raw": source_row.get("gender") or "",
        "grade_raw": source_row.get("grade") or "",
        "shall_we_begin": source_row.get("shall_we_begin") or "",
    }

    if not _contains_agree(base["shall_we_begin"]):
        return _prepared_error(base, STATUS_NOT_AGREE, status="Skipped")

    phone = _canonicalize_phone(base["contact_phone_number"])
    if not phone:
        return _prepared_error(base, "Invalid contact_phone_number")
    base["phone"] = phone

    if phone in done_phones:
        return _prepared_error(base, DUPLICATE_DONE_MESSAGE)
    if phone in ready_phones:
        return _prepared_error(base, DUPLICATE_RUN_MESSAGE)

    grade = _normalize_grade(base["grade_raw"])
    if not grade:
        return _prepared_error(base, f"Invalid grade: {base['grade_raw']}")
    base["grade"] = grade
    base["level"] = _level_for_grade(grade)
    base["gender"] = _normalize_gender(base["gender_raw"])
    base["student_name"] = _clean_student_name(base["student_name_raw"])

    school_id, school_error = _get_school_id_for_registration(
        phone,
        base["student_name"],
        school_lookup=school_lookup,
    )
    if school_error:
        return _prepared_error(base, school_error)
    if not frappe.db.exists("School", school_id):
        return _prepared_error(base, f"School not found: {school_id}")
    base["school_id"] = school_id

    school_enrollment, enrollment_error = _get_school_enrollment_for_registration(
        school_id,
        base["timestamp"],
    )
    if enrollment_error:
        return _prepared_error(base, enrollment_error)
    if not school_enrollment:
        return _prepared_error(base, "School Batch Enrollment not found")
    batch = str(school_enrollment.batch_number or "").strip()
    if not batch:
        return _prepared_error(base, "Selected School Batch Enrollment has no batch")
    if not frappe.db.exists("Batch", batch):
        return _prepared_error(base, f"Batch not found: {batch}")
    base["batch"] = batch

    course_result = _get_course_from_school_enrollment(school_enrollment, grade)
    if course_result.get("error"):
        return _prepared_error(base, course_result["error"])
    base["course_names"] = course_result["course_names"]
    base["course_vertical"] = course_result["course_vertical"]
    base["prepare_status"] = "Ready"
    base["message"] = ""
    return base


def _upsert_student(row: dict, import_user: str):
    existing_name = _find_existing_student(row["phone"], row["student_name"])
    if existing_name:
        student = frappe.get_doc("Student", existing_name)
        student.name1 = row["student_name"]
        student.phone = row["phone"]
        student.school_id = row["school_id"]
        student.grade = row["grade"]
        student.language = row["language_id"]
        student.whatsapp_consent = 1
        if row.get("gender"):
            student.gender = row["gender"]
    else:
        student = frappe.get_doc({
            "doctype": "Student",
            "name1": row["student_name"],
            "phone": row["phone"],
            "gender": row.get("gender") or None,
            "school_id": row["school_id"],
            "grade": row["grade"],
            "language": row["language_id"],
            "whatsapp_consent": 1,
            "joined_on": frappe.utils.now_datetime().date(),
            "status": "active",
        })

    _append_enrollment_if_missing(student, row)

    if existing_name:
        student.modified_by = import_user
        student.save(ignore_permissions=True)
        return student

    student.owner = import_user
    student.modified_by = import_user
    return _insert_student(student)


def _insert_student(student):
    student.name = _reserve_next_student_name()
    old_in_import = getattr(frappe.flags, "in_import", False)
    frappe.flags.in_import = True
    try:
        student.insert(ignore_permissions=True)
    finally:
        frappe.flags.in_import = old_in_import
    return student


def _append_enrollment_if_missing(student, row: dict) -> None:
    batch = str(row.get("batch") or "").strip()
    if not batch:
        return

    existing_batches = {
        str(enrollment.batch or "").strip()
        for enrollment in (student.get("enrollment") or [])
    }
    if batch in existing_batches:
        return

    student.append("enrollment", {
        "batch": batch,
        "vertical": row.get("course_vertical") or "",
        "level": row.get("level") or "",
        "grade": row.get("grade") or "",
        "date_joining": frappe.utils.now_datetime().date(),
        "school": row.get("school_id") or "",
        "whatsapp_response": 0,
    })


def _find_existing_student(phone: str, student_name: str) -> str | None:
    variants = _phone_variants(phone)
    rows = frappe.db.sql("""
        SELECT
            s.name,
            s.name1,
            max(se.date_joining) AS latest_date_joining
        FROM "tabStudent" s
        LEFT JOIN "tabStudent Enrollment" se
          ON se.parent = s.name
        WHERE s.phone IN %(phones)s
        GROUP BY s.name, s.name1, s.modified
        ORDER BY
            CASE
                WHEN lower(coalesce(s.name1, '')) = lower(%(student_name)s)
                THEN 1 ELSE 0
            END DESC,
            latest_date_joining DESC NULLS LAST,
            s.modified DESC,
            s.name ASC
        LIMIT 1
    """, {
        "phones": tuple(variants),
        "student_name": student_name,
    }, as_dict=True)
    return rows[0]["name"] if rows else None


def _read_all_source_sheets(session: AuthorizedSession, ensure_status_column: bool) -> list[SourceSheet]:
    sheets = []
    for config in FIXED_STUDENT_REGISTRATION_SHEETS:
        sheets.append(_read_source_sheet(session, config, ensure_status_column=ensure_status_column))
    return sheets


def _read_source_sheet(session: AuthorizedSession, config: dict, ensure_status_column: bool) -> SourceSheet:
    metadata = _sheets_get(
        session,
        f"https://sheets.googleapis.com/v4/spreadsheets/{config['spreadsheet_id']}",
        params={
            "fields": "properties.title,sheets(properties(sheetId,title,index))",
        },
    )
    sheet_meta = _select_sheet_metadata(metadata, config["sheet_id"])
    sheet_title = sheet_meta["title"]
    values = _get_sheet_values(session, config["spreadsheet_id"], sheet_title)
    if not values:
        raise frappe.ValidationError(f"Sheet is empty: {config['spreadsheet_id']}")

    header = [str(value or "").strip() for value in values[0]]
    header_map = _resolve_source_header_map(header)
    status_column_index = _find_header_index(header, REGISTRATION_STATUS_COLUMN)
    if status_column_index is None:
        if not ensure_status_column:
            raise frappe.ValidationError("registration_status column is missing")
        status_column_index = len(header) + 1
        _write_status_updates(session, [{
            "spreadsheet_id": config["spreadsheet_id"],
            "range": _cell_range(sheet_title, status_column_index, 1),
            "value": REGISTRATION_STATUS_COLUMN,
        }])
        header.append(REGISTRATION_STATUS_COLUMN)

    process_status_column_index = _find_header_index(header, PROCESS_STATUS_COLUMN)
    if process_status_column_index is None:
        if not ensure_status_column:
            raise frappe.ValidationError(f"{PROCESS_STATUS_COLUMN} column is missing")
        process_status_column_index = len(header) + 1
        _write_status_updates(session, [{
            "spreadsheet_id": config["spreadsheet_id"],
            "range": _cell_range(sheet_title, process_status_column_index, 1),
            "value": PROCESS_STATUS_COLUMN,
        }])
        header.append(PROCESS_STATUS_COLUMN)

    rows = []
    for row_number, raw_row in enumerate(values[1:], start=2):
        row = {
            key: _get_raw_cell(raw_row, index)
            for key, index in header_map.items()
        }
        row["registration_status"] = _get_raw_cell(raw_row, status_column_index - 1)
        row["process_status"] = _get_raw_cell(raw_row, process_status_column_index - 1)
        row["status_range"] = _cell_range(sheet_title, status_column_index, row_number)
        row["process_status_range"] = _cell_range(sheet_title, process_status_column_index, row_number)
        row["row_number"] = row_number
        rows.append(row)

    return SourceSheet(
        language=config["language"],
        spreadsheet_id=config["spreadsheet_id"],
        sheet_id=int(sheet_meta["sheetId"]),
        spreadsheet_title=metadata.get("properties", {}).get("title") or "",
        sheet_title=sheet_title,
        header_map=header_map,
        status_column_index=status_column_index,
        process_status_column_index=process_status_column_index,
        rows=rows,
    )


def _resolve_source_header_map(header: list[str]) -> dict[str, int]:
    header_indexes = {
        _normalize_header(value): index
        for index, value in enumerate(header)
        if str(value or "").strip()
    }
    missing = [column for column in SOURCE_COLUMNS if column not in header_indexes]
    if missing:
        raise frappe.ValidationError(f"Source sheet is missing required columns: {missing}")
    return {
        target_key: header_indexes[column_name]
        for column_name, target_key in SOURCE_COLUMNS.items()
    }


def _select_sheet_metadata(metadata: dict, preferred_sheet_id: int) -> dict:
    sheets = metadata.get("sheets") or []
    for sheet in sheets:
        props = sheet.get("properties") or {}
        if int(props.get("sheetId")) == int(preferred_sheet_id):
            return props
    if not sheets:
        raise frappe.ValidationError("Spreadsheet has no sheets")
    return sorted(
        [sheet.get("properties") or {} for sheet in sheets],
        key=lambda props: int(props.get("index") or 0),
    )[0]


def _get_sheet_values(session: AuthorizedSession, spreadsheet_id: str, sheet_title: str) -> list[list[str]]:
    range_name = _quote_sheet_name(sheet_title)
    data = _sheets_get(
        session,
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values/{quote(range_name, safe='')}",
        params={"majorDimension": "ROWS"},
    )
    return data.get("values") or []


def _get_sheets_session() -> AuthorizedSession:
    credentials = _get_google_service_account_credentials(
        GCP_CREDENTIALS_PROJECT_ID,
        scopes=GOOGLE_SHEETS_READWRITE_SCOPES,
    )
    return AuthorizedSession(credentials)


def _sheets_get(session: AuthorizedSession, url: str, params: dict | None = None) -> dict:
    response = session.get(url, params=params or {}, timeout=120)
    _raise_for_sheets_response(response)
    return response.json()


def _sheets_post(session: AuthorizedSession, url: str, body: dict) -> dict:
    response = session.post(url, json=body, timeout=120)
    _raise_for_sheets_response(response)
    return response.json() if response.content else {}


def _raise_for_sheets_response(response) -> None:
    if response.ok:
        return
    try:
        detail = response.json()
    except Exception:
        detail = response.text
    raise frappe.ValidationError(f"Google Sheets API failed: {response.status_code} {detail}")


def _write_status_updates(session: AuthorizedSession, updates: list[dict]) -> None:
    updates_by_sheet: dict[str, list[dict]] = {}
    for update in updates:
        value = str(update.get("value") or "")
        updates_by_sheet.setdefault(update["spreadsheet_id"], []).append({
            "range": update["range"],
            "values": [[value[:1000]]],
        })

    for spreadsheet_id, data in updates_by_sheet.items():
        for chunk in _chunks(data, 500):
            _sheets_post(
                session,
                f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values:batchUpdate",
                {
                    "valueInputOption": "RAW",
                    "data": chunk,
                },
            )


def _get_language_id(language_name: str) -> str:
    return frappe.db.get_value("TAP Language", {"language_name": language_name}, "name") or ""


def _get_latest_student_consent(phone: str) -> dict | None:
    rows = frappe.db.sql("""
        SELECT name, phone_number, school, whatsapp_consent, timestamp
        FROM "tabStudent Consent"
        WHERE phone_number IN %(phones)s
        ORDER BY timestamp DESC NULLS LAST, modified DESC
        LIMIT 1
    """, {"phones": tuple(_phone_variants(phone))}, as_dict=True)
    return rows[0] if rows else None


def _get_school_id_for_registration(
    phone: str,
    _student_name: str,
    school_lookup: "_GlificSchoolLookup | None" = None,
) -> tuple[str, str]:
    consent = _get_latest_student_consent(phone)
    if consent:
        school_id = str(consent.get("school") or "").strip()
        if not school_id:
            return "", "Student Consent has no school"
        return school_id, ""

    if school_lookup:
        glific_school_id, glific_error = school_lookup.get_school_id(phone)
    else:
        glific_school_id, glific_error = _get_school_id_from_glific(phone)
    if glific_school_id:
        return glific_school_id, ""
    if glific_error:
        return "", glific_error

    return "", "Student Consent not found and Glific contact school_id not found"


class _GlificSchoolLookup:
    def __init__(self) -> None:
        self._contact_school_cache: dict[str, tuple[str, str]] = {}

    def get_school_id(self, phone: str) -> tuple[str, str]:
        phone = _canonicalize_phone(phone) or ""
        if not phone:
            return "", ""

        if phone not in self._contact_school_cache:
            try:
                self._contact_school_cache[phone] = _get_school_id_from_glific_contact(phone)
            except Exception as exc:
                self._contact_school_cache[phone] = (
                    "",
                    f"Glific contact school lookup failed: {_short_error(exc)}",
                )
        return self._contact_school_cache[phone]


def _get_school_id_from_glific(phone: str) -> tuple[str, str]:
    lookup = _GlificSchoolLookup()
    return lookup.get_school_id(phone)


def _get_school_id_from_glific_contact(phone: str) -> tuple[str, str]:
    from tap_lms.glific_integration import get_contact_by_phone

    for candidate_phone in _phone_variants(phone):
        contact = get_contact_by_phone(candidate_phone)
        if not contact:
            continue

        fields = _parse_json_dict(contact.get("fields"))
        school_id = _extract_glific_contact_field(fields, GLIFIC_CONTACT_FIELD_SCHOOL_ID)
        if school_id:
            return school_id, ""

    return "", ""


def _extract_glific_contact_field(fields: dict, fieldname: str) -> str:
    value = fields.get(fieldname)
    if isinstance(value, dict):
        value = value.get("value")
    return str(value or "").strip()


def _parse_json_dict(value: object) -> dict:
    parsed = value
    for _ in range(3):
        if isinstance(parsed, dict):
            return parsed
        if not isinstance(parsed, str):
            return {}
        try:
            parsed = json.loads(parsed)
        except json.JSONDecodeError:
            return {}
    return parsed if isinstance(parsed, dict) else {}


def _get_school_enrollment_for_registration(
    school_id: str,
    registration_timestamp: object,
):
    registration_dt = _parse_registration_timestamp(registration_timestamp)
    if not registration_dt:
        return None, f"Invalid registration timestamp: {registration_timestamp}"

    school = frappe.get_doc("School", school_id)
    candidates = []
    for enrollment in school.get("batch_enrollments") or []:
        enrollment_dt = _parse_enrollment_timestamp(getattr(enrollment, "doj", None))
        if not enrollment_dt or enrollment_dt > registration_dt:
            continue
        candidates.append((enrollment_dt, int(getattr(enrollment, "idx", 0) or 0), enrollment))

    if not candidates:
        return None, (
            "School Batch Enrollment not found on or before "
            f"registration timestamp {registration_timestamp}"
        )

    return max(candidates, key=lambda item: (item[0], item[1]))[2], ""


def _parse_registration_timestamp(value: object) -> datetime | None:
    return _parse_datetime_value(value, date_only_time=time.max)


def _parse_enrollment_timestamp(value: object) -> datetime | None:
    return _parse_datetime_value(value, date_only_time=time.min)


def _parse_datetime_value(value: object, date_only_time: time) -> datetime | None:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, date_only_time)
    else:
        raw = str(value or "").strip()
        if not raw:
            return None
        try:
            parsed = date_parser.parse(raw, fuzzy=True, dayfirst=_should_parse_day_first(raw))
        except (TypeError, ValueError, OverflowError):
            return None
        if (
            parsed.hour == 0
            and parsed.minute == 0
            and parsed.second == 0
            and parsed.microsecond == 0
            and not re.search(r"\d{1,2}:\d{2}", raw)
        ):
            parsed = datetime.combine(parsed.date(), date_only_time)

    if parsed.tzinfo:
        parsed = parsed.astimezone().replace(tzinfo=None)
    return parsed


def _should_parse_day_first(raw: str) -> bool:
    match = re.match(r"^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", raw)
    if not match:
        return False
    first = int(match.group(1))
    second = int(match.group(2))
    if first > 12:
        return True
    if second > 12:
        return False
    return False


def _get_course_from_school_enrollment(school_enrollment, grade: str) -> dict:
    grades_courses = getattr(school_enrollment, "grades_courses", None)
    if not grades_courses:
        return {"error": "Selected School Batch Enrollment has no grades_courses"}
    if not isinstance(grades_courses, dict):
        try:
            grades_courses = frappe.parse_json(grades_courses)
        except Exception:
            return {"error": "Selected School Batch Enrollment grades_courses is invalid JSON"}
    if not isinstance(grades_courses, dict):
        return {"error": "Selected School Batch Enrollment grades_courses must be a JSON object"}

    value = grades_courses.get(str(grade))
    if isinstance(value, str) and value.strip():
        course_names = [value.strip()]
    elif isinstance(value, list):
        course_names = [str(item).strip() for item in value if str(item or "").strip()]
    else:
        course_names = []

    if not course_names:
        return {"error": f"Course mapping not found for grade {grade}"}
    if len(course_names) > 1:
        return {"course_names": course_names, "course_vertical": ""}

    course_vertical = frappe.db.get_value(
        "Course Verticals",
        {"name2": course_names[0]},
        "name",
    ) or ""
    if not course_vertical:
        return {"error": f"Course vertical not found: {course_names[0]}"}
    return {"course_names": course_names, "course_vertical": course_vertical}


def _collect_done_phones(sheets: Iterable[SourceSheet]) -> set[str]:
    done_phones: set[str] = set()
    for sheet in sheets:
        for row in sheet.rows:
            if not _status_is_done(row.get("registration_status")):
                continue
            phone = _canonicalize_phone(row.get("contact_phone_number"))
            if phone:
                done_phones.add(phone)
    return done_phones


def _index_current_rows(sheets: Iterable[SourceSheet]) -> dict[tuple[str, int, int], dict]:
    index = {}
    for sheet in sheets:
        for row in sheet.rows:
            index[(sheet.spreadsheet_id, sheet.sheet_id, row["row_number"])] = row
    return index


def _source_row_key(row: dict) -> tuple[str, int, int]:
    return (row["spreadsheet_id"], int(row["sheet_id"]), int(row["row_number"]))


def _prepared_error(base: dict, message: str, status: str = "Error") -> dict:
    row = dict(base)
    row.setdefault("phone", "")
    row.setdefault("student_name", _clean_student_name(row.get("student_name_raw")))
    row.setdefault("gender", _normalize_gender(row.get("gender_raw")))
    row.setdefault("grade", _normalize_grade(row.get("grade_raw")) or "")
    row.setdefault("school_id", "")
    row.setdefault("batch", "")
    row.setdefault("course_names", [])
    row.setdefault("course_vertical", "")
    row.setdefault("level", _level_for_grade(row.get("grade")) if row.get("grade") else "")
    row["prepare_status"] = status
    row["message"] = message
    row["final_registration_status"] = message
    row["final_process_status"] = _process_status_for_row(row)
    return row


def _failed_copy(row: dict, message: str) -> dict:
    failed = dict(row)
    failed["prepare_status"] = "Error"
    failed["message"] = message
    failed["final_registration_status"] = message
    failed["final_process_status"] = _process_status_for_row(failed)
    return failed


def _status_update(row: dict, value: str) -> dict:
    return {
        "spreadsheet_id": row["spreadsheet_id"],
        "range": row["status_range"],
        "value": value,
    }


def _process_status_update(row: dict, value: str) -> dict:
    range_name = row.get("process_status_range")
    if not range_name:
        column_index = int(row.get("process_status_column_index") or 0)
        if not column_index and row.get("status_column_index"):
            column_index = int(row["status_column_index"]) + 1
        range_name = _cell_range(row["sheet_title"], column_index, row["row_number"])

    return {
        "spreadsheet_id": row["spreadsheet_id"],
        "range": range_name,
        "value": value,
    }


def _process_status_for_row(row: dict) -> str:
    message = str(row.get("message") or "")
    if message.startswith("Duplicate contact_phone_number"):
        return PROCESS_STATUS_COMPLETE
    return PROCESS_STATUS_FAIL


def _process_status_update_from_source(sheet: SourceSheet, source_row: dict, value: str) -> dict:
    return {
        "spreadsheet_id": sheet.spreadsheet_id,
        "range": _cell_range(sheet.sheet_title, sheet.process_status_column_index, source_row["row_number"]),
        "value": value,
    }


def _create_and_upload_prepared_workbook(rows: list[dict]) -> str:
    return _create_and_upload_rows_workbook(
        rows,
        prefix="student-sheet-registration/prepared",
        filename_prefix="student_sheet_registration_prepared",
    )


def _create_and_upload_failed_rows_workbook(rows: list[dict]) -> str:
    return _create_and_upload_rows_workbook(
        rows,
        prefix="student-sheet-registration/failures",
        filename_prefix="student_sheet_registration_failures",
    )


def _create_and_upload_not_done_rows_csv(rows: list[dict]) -> str:
    timestamp = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
    object_name = (
        "student-sheet-registration/not-done/"
        f"student_sheet_registration_not_done_{timestamp}.csv"
    )
    return _upload_bytes_to_gcs(
        _render_not_done_rows_csv(rows),
        object_name,
        FAILED_ROWS_GCP_PROJECT_ID,
        content_type="text/csv",
    )


def _render_not_done_rows_csv(rows: list[dict]) -> bytes:
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=NOT_DONE_ROWS_CSV_HEADERS)
    writer.writeheader()
    writer.writerows(_not_done_csv_row(row) for row in rows)
    return buffer.getvalue().encode("utf-8")


def _not_done_csv_row(row: dict) -> dict[str, str]:
    values = _workbook_row(row) + [
        _not_done_registration_status(row),
        _not_done_process_status(row),
    ]
    return {
        header: str(value or "")
        for header, value in zip(NOT_DONE_ROWS_CSV_HEADERS, values)
    }


def _not_done_registration_status(row: dict) -> str:
    explicit_status = str(
        row.get("final_registration_status")
        or row.get("registration_status")
        or ""
    ).strip()
    if explicit_status:
        return explicit_status
    return str(row.get("message") or row.get("prepare_status") or "").strip()


def _not_done_process_status(row: dict) -> str:
    explicit_status = str(
        row.get("final_process_status")
        or row.get("process_status")
        or ""
    ).strip()
    if explicit_status:
        return explicit_status
    return _process_status_for_row(row)


def _create_and_upload_rows_workbook(rows: list[dict], prefix: str, filename_prefix: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rows"
    ws.append(PREPARED_HEADERS)
    highlight = PatternFill(fill_type="solid", fgColor="FFF59D")

    for row in rows:
        ws.append(_workbook_row(row))
        if row.get("message"):
            excel_row = ws.max_row
            ws.cell(row=excel_row, column=len(PREPARED_HEADERS)).fill = highlight

    buffer = BytesIO()
    wb.save(buffer)
    timestamp = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
    object_name = f"{prefix}/{filename_prefix}_{timestamp}.xlsx"
    return _upload_bytes_to_gcs(buffer.getvalue(), object_name, FAILED_ROWS_GCP_PROJECT_ID)


def _workbook_row(row: dict) -> list[str]:
    return [
        row.get("language") or "",
        row.get("spreadsheet_title") or "",
        row.get("sheet_title") or "",
        str(row.get("row_number") or ""),
        row.get("timestamp") or "",
        row.get("student_name") or row.get("student_name_raw") or "",
        row.get("contact_phone_number") or "",
        row.get("phone") or "",
        row.get("gender") or "",
        row.get("grade") or "",
        row.get("school_id") or "",
        row.get("batch") or "",
        row.get("course_vertical") or "",
        ", ".join(row.get("course_names") or []),
        row.get("level") or "",
        row.get("prepare_status") or "",
        row.get("message") or "",
    ]


def _create_and_upload_glific_contact_csvs(rows: list[dict]) -> list[dict]:
    timestamp = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
    contact_rows = [_glific_contact_row(row) for row in rows]
    file_name = f"student_sheet_registration_contacts_{timestamp}.csv"
    file_path = _upload_glific_contact_csv(file_name, contact_rows)
    return [{
        "tab_name": "All Uploaded Rows",
        "file_name": file_name,
        "file_path": file_path,
        "row_count": len(contact_rows),
    }]


def _glific_contact_row(row: dict) -> dict:
    school_id = row.get("school_id") or ""
    batch = row.get("batch") or ""
    language = row.get("language") or ""
    course_vertical = row.get("course_vertical") or ""
    state_id = frappe.db.get_value("School", school_id, "state") if school_id else ""
    state_name = frappe.db.get_value("State", state_id, "state_name") if state_id else ""
    model_id = frappe.db.get_value("School", school_id, "model") if school_id else ""
    model_name = frappe.db.get_value("Tap Models", model_id, "mname") if model_id else ""
    batch_id = frappe.db.get_value("Batch", batch, "batch_id") if batch else ""
    course = (
        frappe.db.get_value("Course Verticals", course_vertical, "name2")
        if course_vertical else ""
    )
    data = {
        "name": row.get("student_name") or "",
        "phone": row.get("phone") or "",
        "language": language,
        "delete": "0",
        "school_id": school_id,
        "state": state_name or "",
        "model": model_name or "",
        "buddy_name": row.get("student_name") or "",
        "batch_id": batch_id or "",
        "grade": row.get("grade") or "",
        "level": row.get("level") or "",
        "course": course or "",
    }
    return {header: str(data.get(header) or "") for header in GLIFIC_CSV_HEADERS}


def _canonicalize_phone(value: object) -> str | None:
    raw = str(value or "").strip()
    if raw.endswith(".0"):
        raw = raw[:-2]
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 10:
        return f"91{digits}"
    if len(digits) == 12 and digits.startswith("91"):
        return digits
    return None


def _phone_variants(phone: str) -> list[str]:
    phone = _canonicalize_phone(phone) or ""
    if not phone:
        return []
    return [phone, phone[2:]]


def _normalize_grade(value: object) -> str:
    raw = str(value or "").strip()
    if raw.endswith(".0"):
        raw = raw[:-2]
    if not raw.isdigit():
        return ""
    grade = int(raw)
    if 1 <= grade <= 12:
        return str(grade)
    return ""


def _level_for_grade(grade: object) -> str:
    try:
        grade_num = int(str(grade).strip())
    except Exception:
        return ""
    if grade_num <= 3:
        return "Level 0"
    if 4 <= grade_num <= 5:
        return "Level 1"
    if 6 <= grade_num <= 8:
        return "Level 2"
    if 9 <= grade_num <= 10:
        return "Level 3"
    if 11 <= grade_num <= 12:
        return "Level 4"
    return ""


def _normalize_gender(value: object) -> str:
    raw = str(value or "").strip().lower()
    if not raw:
        return ""
    if raw in {"f", "female"} or "girl" in raw:
        return "Female"
    if raw in {"m", "male"} or "boy" in raw:
        return "Male"
    if "other" in raw:
        return "Others"
    if raw in {"na", "n/a", "not available"}:
        return "Not Available"
    return ""


def _clean_student_name(value: object) -> str:
    cleaned = re.sub(r"\s+", " ", str(value or "").strip())
    cleaned = re.sub(r"[^A-Za-z ]+", "", cleaned).strip()
    return cleaned or "Champ"


def _contains_agree(value: object) -> bool:
    return "agree" in str(value or "").lower()


def _status_is_done(value: object) -> bool:
    return str(value or "").strip().lower() == STATUS_DONE.lower()


def _row_is_blank(row: dict) -> bool:
    return not any(str(row.get(key) or "").strip() for key in SOURCE_COLUMNS.values())


def _normalize_header(value: object) -> str:
    return re.sub(r"[\s-]+", "_", str(value or "").strip().lower())


def _find_header_index(header: list[str], column_name: str) -> int | None:
    target = _normalize_header(column_name)
    for index, value in enumerate(header, start=1):
        if _normalize_header(value) == target:
            return index
    return None


def _get_raw_cell(row: list, index: int) -> str:
    if index >= len(row):
        return ""
    return str(row[index] or "").strip()


def _quote_sheet_name(sheet_title: str) -> str:
    return "'" + str(sheet_title).replace("'", "''") + "'"


def _cell_range(sheet_title: str, column_index: int, row_number: int) -> str:
    return f"{_quote_sheet_name(sheet_title)}!{_column_letter(column_index)}{row_number}"


def _column_letter(index: int) -> str:
    letters = []
    while index:
        index, remainder = divmod(index - 1, 26)
        letters.append(chr(65 + remainder))
    return "".join(reversed(letters))


def _chunks(values: list, size: int):
    for start in range(0, len(values), size):
        yield values[start:start + size]


def _short_error(exc: Exception) -> str:
    message = str(exc or "").strip() or type(exc).__name__
    return message[:1000]


def _emit(log_fn, message: str) -> None:
    if log_fn:
        log_fn(message)
    else:
        print(message)


def _emit_progress(progress_fn, payload: dict) -> None:
    if progress_fn:
        progress_fn(payload)
