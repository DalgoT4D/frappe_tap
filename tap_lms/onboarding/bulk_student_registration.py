from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Callable, Iterable
from urllib.parse import urlparse

import frappe
import requests
from google.auth.transport.requests import AuthorizedSession
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from psycopg2.extras import execute_values

from tap_lms.onboarding.backend_upload_utils import (
    FAILED_ROWS_GCP_PROJECT_ID,
    GCP_CREDENTIALS_PROJECT_ID,
    GLIFIC_EXISTING_STUDENT_CSV_HEADERS,
    GLIFIC_NEW_STUDENT_CSV_HEADERS,
    get_google_service_account_credentials as _get_google_service_account_credentials,
    get_google_service_account_email as _get_google_service_account_email,
    upload_bytes_to_gcs as _upload_bytes_to_gcs,
    upload_glific_contact_csv as _upload_glific_contact_csv,
)


# Bench console usage:
# from tap_lms.onboarding.bulk_student_registration import run_import
# run_import()
#
# Or override at call time:
# run_import(
#     spreadsheet_url="https://docs.google.com/spreadsheets/d/18VmTVTCwK2QNct3Y-aICdH7I6F2Zq8ltSWIIqIlRtgg/edit?gid=0#gid=0",
#     tab_names=["Existing Students", "New Students"],
#     sample_test=100,
# )

SPREADSHEET_URL = "https://docs.google.com/spreadsheets/d/18VmTVTCwK2QNct3Y-aICdH7I6F2Zq8ltSWIIqIlRtgg/edit?gid=0#gid=0"
TAB_NAMES = ["Existing Students", "New Students"]
SAMPLE_TEST = 0
BATCH_SIZE = 100
IMPORT_USER = "Administrator"
EXPECTED_COLUMNS = {
    "Student Name": "student_name_raw",
    "Contact No.": "contact_no_raw",
    "Gender": "gender_raw",
    "School ID": "school_id_raw",
    "Language": "language_raw",
    "Batch": "batch_raw",
    "Grade": "grade_raw",
    "Course": "course_raw",
}

FAILED_WORKBOOK_HEADERS = [
    "Student Name",
    "Contact No.",
    "Gender",
    "Batch",
    "Course",
    "Grade",
    "School ID",
    "Language",
]
REMAINING_ROWS_WORKBOOK_HEADERS = [
    "Student Name",
    "Contact No.",
    "Gender",
    "School ID",
    "Language",
    "Batch",
    "Grade",
    "Course",
]

@dataclass(frozen=True)
class RawRow:
    source_tab: str
    source_priority: int
    row_in_tab: int
    student_name_raw: str
    contact_no_raw: str
    gender_raw: str
    school_id_raw: str
    language_raw: str
    batch_raw: str
    grade_raw: str
    course_raw: str


def _emit(log_fn: Callable[[str], None] | None, message: str) -> None:
    if log_fn:
        log_fn(message)
    else:
        print(message)


def _emit_progress(
    progress_fn: Callable[[dict], None] | None,
    payload: dict,
) -> None:
    if progress_fn:
        progress_fn(payload)


def _should_stop_before_timeout(
    started_at: datetime,
    timeout_seconds: int | None,
    stop_before_timeout_seconds: int,
) -> bool:
    if not timeout_seconds:
        return False
    deadline = started_at + timedelta(seconds=timeout_seconds)
    remaining_seconds = (deadline - frappe.utils.now_datetime()).total_seconds()
    return remaining_seconds <= stop_before_timeout_seconds


def run_import(
    spreadsheet_url: str | None = None,
    tab_names: list[str] | None = None,
    sample_test: int | None = None,
    batch_size: int | None = None,
    import_user: str | None = None,
    job_name: str | None = None,
    import_date: date | str | None = None,
    timeout_seconds: int | None = None,
    stop_before_timeout_seconds: int = 600,
    log_fn: Callable[[str], None] | None = None,
    progress_fn: Callable[[dict], None] | None = None,
) -> dict:
    spreadsheet_url = (spreadsheet_url or SPREADSHEET_URL or "").strip()
    tab_names = list(tab_names or TAB_NAMES or [])
    sample_test = int(SAMPLE_TEST if sample_test is None else sample_test)
    batch_size = int(BATCH_SIZE if batch_size is None else batch_size)
    import_user = (import_user or IMPORT_USER or frappe.session.user or "Administrator").strip()

    if not spreadsheet_url:
        raise ValueError("SPREADSHEET_URL is required")
    if not tab_names:
        raise ValueError("TAB_NAMES must contain at least one sheet name")
    if sample_test < 0:
        raise ValueError("sample_test must be >= 0")
    if batch_size <= 0:
        raise ValueError("batch_size must be > 0")

    import_date = frappe.utils.getdate(import_date) if import_date else date.today()
    started_at = frappe.utils.now_datetime()
    _emit(
        log_fn,
        f"[student-import] start import_date={import_date} "
        f"tabs={tab_names} sample_test={sample_test} batch_size={batch_size}"
    )
    try:
        workbook = _download_workbook(spreadsheet_url)
        _prepare_temp_tables()
        raw_rows = _stage_selected_tabs(workbook, tab_names, log_fn=log_fn)
        if not raw_rows:
            raise ValueError("No rows found in selected tabs")
        _emit(log_fn, f"[student-import] loaded raw rows={raw_rows}")
        _build_clean_stage(sample_test)
        _build_validation_tables(sample_test)

        precheck = _run_prechecks()
        _print_precheck(precheck, log_fn=log_fn)
        failed_rows_file_url = ""
        if precheck["invalid_rows"] > 0:
            try:
                failed_rows_file_url = _create_and_upload_failed_rows_workbook(tab_names, job_name=job_name)
                _emit(log_fn, f"[student-import] failed_rows_file_url={failed_rows_file_url}")
            except Exception as exc:
                _emit(log_fn, f"[student-import] failed_rows_workbook_upload_failed error={exc}")

        total_effective_rows = precheck["valid_rows"]
        summary = {
            "updated_students": 0,
            "inserted_students": 0,
            "inserted_enrollments": 0,
            "skipped_duplicate_enrollments": 0,
            "batches_processed": 0,
            "effective_rows": total_effective_rows,
            "processed_rows": 0,
            "remaining_rows_count": 0,
            "remaining_rows": "",
            "stopped_before_timeout": False,
            "failed_rows": precheck["invalid_rows"],
            "skipped_rows": precheck["invalid_rows"],
            "failed_rows_file_url": failed_rows_file_url,
            "glific_contact_files": [],
        }
        for batch_no, offset in enumerate(range(0, total_effective_rows, batch_size), start=1):
            _prepare_batch_subset(offset=offset, batch_size=batch_size)
            batch_summary = _execute_import(
                import_date=import_date,
                import_user=import_user,
                source_table="tmp_student_import_batch",
            )
            frappe.db.commit()
            summary["updated_students"] += batch_summary["updated_students"]
            summary["inserted_students"] += batch_summary["inserted_students"]
            summary["inserted_enrollments"] += batch_summary["inserted_enrollments"]
            summary["skipped_duplicate_enrollments"] += batch_summary["skipped_duplicate_enrollments"]
            summary["batches_processed"] = batch_no
            summary["processed_rows"] = min(offset + batch_summary["batch_rows"], total_effective_rows)
            elapsed = frappe.utils.now_datetime() - started_at
            batch_message = (
                "[student-import] batch completed "
                f"batch_no={batch_no} "
                f"batch_rows={batch_summary['batch_rows']} "
                f"updated_in_batch={batch_summary['updated_students']} "
                f"inserted_in_batch={batch_summary['inserted_students']} "
                f"enrollments_in_batch={batch_summary['inserted_enrollments']} "
                f"duplicate_enrollments_skipped={batch_summary['skipped_duplicate_enrollments']} "
                f"processed_total={summary['processed_rows']}/{total_effective_rows} "
                f"elapsed={elapsed}"
            )
            _emit(log_fn, batch_message)
            _emit_progress(progress_fn, {
                "event": "batch_completed",
                "batch_no": batch_no,
                "batch_rows": batch_summary["batch_rows"],
                "processed_total": summary["processed_rows"],
                "effective_rows": total_effective_rows,
                "elapsed": str(elapsed),
                "summary": dict(summary),
                "message": batch_message,
            })
            if (
                summary["processed_rows"] < total_effective_rows
                and _should_stop_before_timeout(started_at, timeout_seconds, stop_before_timeout_seconds)
            ):
                remaining_rows_url, remaining_rows_count = _create_and_upload_remaining_rows_workbook(
                    processed_total=int(summary["processed_rows"]),
                    tab_names=tab_names,
                    job_name=job_name,
                )
                summary["remaining_rows"] = remaining_rows_url
                summary["remaining_rows_count"] = remaining_rows_count
                summary["stopped_before_timeout"] = True
                summary["stop_reason"] = (
                    f"Stopped with <= {stop_before_timeout_seconds} seconds remaining before worker timeout."
                )
                stop_message = (
                    "[student-import] stopped before timeout "
                    f"processed_total={summary['processed_rows']}/{total_effective_rows} "
                    f"remaining_rows={remaining_rows_count} "
                    f"remaining_rows_file_url={remaining_rows_url}"
                )
                _emit(log_fn, stop_message)
                _emit_progress(progress_fn, {
                    "event": "stopped_before_timeout",
                    "processed_total": summary["processed_rows"],
                    "effective_rows": total_effective_rows,
                    "elapsed": str(frappe.utils.now_datetime() - started_at),
                    "summary": dict(summary),
                    "message": stop_message,
                })
                frappe.db.commit()
                break

        try:
            summary["glific_contact_files"] = _create_and_upload_glific_contact_csvs(tab_names)
            for export_file in summary["glific_contact_files"]:
                _emit(
                    log_fn,
                    "[student-import] glific_contacts_file "
                    f"tab={export_file['tab_name']} "
                    f"student_type={export_file.get('student_type') or ''} "
                    f"rows={export_file['row_count']} "
                    f"url={export_file['file_path']}"
                )
        except Exception as exc:
            summary["glific_contact_files"] = []
            summary["glific_contact_files_error"] = str(exc)
            _emit(log_fn, f"[student-import] glific_contacts_csv_upload_failed error={exc}")

        summary["import_date"] = str(import_date)
        summary["sample_test"] = sample_test
        summary["batch_size"] = batch_size
        summary["tabs"] = tab_names
        summary["elapsed"] = str(frappe.utils.now_datetime() - started_at)
        _emit(log_fn, f"[student-import] success summary={summary}")
        _emit_progress(progress_fn, {
            "event": "completed",
            "summary": dict(summary),
            "message": f"[student-import] success summary={summary}",
        })
        return summary
    except Exception:
        frappe.db.rollback()
        raise


def _download_workbook(spreadsheet_url: str):
    download_url = _build_download_url(spreadsheet_url)
    response = requests.get(download_url, timeout=120)
    if response.ok:
        return load_workbook(BytesIO(response.content), read_only=True, data_only=True)

    file_id = _extract_sheet_id(spreadsheet_url)
    if not file_id:
        response.raise_for_status()

    private_response = _download_private_workbook(file_id)
    return load_workbook(BytesIO(private_response.content), read_only=True, data_only=True)


def _build_download_url(spreadsheet_url: str) -> str:
    sheet_id = _extract_sheet_id(spreadsheet_url)
    if sheet_id:
        return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"

    parsed = urlparse(spreadsheet_url)
    if parsed.scheme in {"http", "https"}:
        return spreadsheet_url
    raise ValueError("Unsupported spreadsheet URL")


def _extract_sheet_id(spreadsheet_url: str) -> str | None:
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", spreadsheet_url)
    return match.group(1) if match else None


def _download_private_workbook(file_id: str) -> requests.Response:
    credentials = _get_google_service_account_credentials()
    session = AuthorizedSession(credentials)
    response = session.get(
        "https://www.googleapis.com/drive/v3/files/{file_id}/export".format(file_id=file_id),
        params={
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
        timeout=120,
    )
    if _is_file_not_exportable_response(response):
        media_response = session.get(
            "https://www.googleapis.com/drive/v3/files/{file_id}".format(file_id=file_id),
            params={"alt": "media"},
            timeout=120,
        )
        if media_response.status_code == 403:
            service_account_email = _get_google_service_account_email()
            raise frappe.ValidationError(
                "Private Google Drive file access denied (403). "
                f"Share the file with the service account '{service_account_email}' "
                f"from GCS Settings project_id '{GCP_CREDENTIALS_PROJECT_ID}', "
                "or verify that the Google Drive API is enabled for that project."
            )
        media_response.raise_for_status()
        return media_response
    if response.status_code == 403:
        service_account_email = _get_google_service_account_email()
        raise frappe.ValidationError(
            "Private Google Sheet access denied (403). "
            f"Share the sheet with the service account '{service_account_email}' "
            f"from GCS Settings project_id '{GCP_CREDENTIALS_PROJECT_ID}', "
            "or verify that the Google Drive API is enabled for that project."
        )
    response.raise_for_status()
    return response


def _is_file_not_exportable_response(response: requests.Response) -> bool:
    if response.status_code != 403:
        return False
    try:
        payload = response.json()
    except Exception:
        return False
    errors = payload.get("error", {}).get("errors", [])
    return any(error.get("reason") == "fileNotExportable" for error in errors if isinstance(error, dict))


def _create_and_upload_failed_rows_workbook(tab_names: list[str], job_name: str | None = None) -> str:
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    header = FAILED_WORKBOOK_HEADERS + ["Validation Errors"]
    highlight = PatternFill(fill_type="solid", fgColor="FFF59D")
    field_to_col = {
        "student_name_raw": 1,
        "contact_no_raw": 2,
        "gender_raw": 3,
        "batch_raw": 4,
        "course_raw": 5,
        "grade_raw": 6,
        "school_id_raw": 7,
        "language_raw": 8,
    }

    invalid_rows = _rows("""
        SELECT
            source_tab,
            row_in_tab,
            student_name_raw,
            contact_no_raw,
            gender_raw,
            school_id_raw,
            language_raw,
            batch_raw,
            grade_raw,
            course_raw,
            validation_errors,
            fail_phone,
            fail_duplicate_phone,
            fail_grade,
            fail_school,
            fail_batch,
            fail_language,
            fail_course,
            fail_ambiguous_phone
        FROM tmp_student_import_invalid
        ORDER BY source_priority, row_in_tab
    """)

    rows_by_tab: dict[str, list[dict]] = {tab_name: [] for tab_name in tab_names}
    for row in invalid_rows:
        rows_by_tab.setdefault(row["source_tab"], []).append(row)

    for tab_name in tab_names:
        ws = wb.create_sheet(title=tab_name[:31] or "Sheet")
        ws.append(header)
        for row in rows_by_tab.get(tab_name, []):
            ws.append([
                row.get("student_name_raw") or "",
                row.get("contact_no_raw") or "",
                row.get("gender_raw") or "",
                row.get("batch_raw") or "",
                row.get("course_raw") or "",
                row.get("grade_raw") or "",
                row.get("school_id_raw") or "",
                row.get("language_raw") or "",
                row.get("validation_errors") or "",
            ])
            excel_row = ws.max_row
            if row.get("fail_phone") or row.get("fail_duplicate_phone") or row.get("fail_ambiguous_phone"):
                ws.cell(row=excel_row, column=field_to_col["contact_no_raw"]).fill = highlight
            if row.get("fail_grade"):
                ws.cell(row=excel_row, column=field_to_col["grade_raw"]).fill = highlight
            if row.get("fail_school"):
                ws.cell(row=excel_row, column=field_to_col["school_id_raw"]).fill = highlight
            if row.get("fail_batch"):
                ws.cell(row=excel_row, column=field_to_col["batch_raw"]).fill = highlight
            if row.get("fail_language"):
                ws.cell(row=excel_row, column=field_to_col["language_raw"]).fill = highlight
            if row.get("fail_course"):
                ws.cell(row=excel_row, column=field_to_col["course_raw"]).fill = highlight
            ws.cell(row=excel_row, column=len(header)).fill = highlight

    buffer = BytesIO()
    wb.save(buffer)
    timestamp = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
    file_stem = _build_failed_rows_filename_stem(job_name, timestamp)
    object_name = f"student-bulk-import-failures/{file_stem}.xlsx"
    return _upload_bytes_to_gcs(buffer.getvalue(), object_name, FAILED_ROWS_GCP_PROJECT_ID)


def _create_and_upload_remaining_rows_workbook(
    processed_total: int,
    tab_names: list[str],
    job_name: str | None = None,
) -> tuple[str, int]:
    remaining_rows = _rows("""
        SELECT
            source_tab,
            student_name_raw AS "Student Name",
            contact_no_raw AS "Contact No.",
            gender_raw AS "Gender",
            school_id_raw AS "School ID",
            language_raw AS "Language",
            batch_raw AS "Batch",
            grade_raw AS "Grade",
            course_raw AS "Course"
        FROM (
            SELECT
                v.*,
                row_number() OVER (ORDER BY source_priority, row_in_tab) AS import_position
            FROM tmp_student_import_valid v
        ) ordered
        WHERE import_position > %s
        ORDER BY import_position
    """, (processed_total,))

    rows_by_tab: dict[str, list[dict]] = {tab_name: [] for tab_name in tab_names}
    for row in remaining_rows:
        rows_by_tab.setdefault(row["source_tab"], []).append(row)

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    for tab_name in tab_names:
        ws = wb.create_sheet(title=tab_name or "Sheet")
        ws.append(REMAINING_ROWS_WORKBOOK_HEADERS)
        for row in rows_by_tab.get(tab_name, []):
            ws.append([
                row.get(header) or ""
                for header in REMAINING_ROWS_WORKBOOK_HEADERS
            ])

    buffer = BytesIO()
    wb.save(buffer)

    timestamp = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
    file_stem = _build_remaining_rows_filename_stem(job_name, timestamp)
    object_name = f"student-bulk-import-remaining/{file_stem}.xlsx"
    file_url = _upload_bytes_to_gcs(
        buffer.getvalue(),
        object_name,
        FAILED_ROWS_GCP_PROJECT_ID,
    )
    return file_url, len(remaining_rows)


def _build_failed_rows_filename_stem(job_name: str | None, timestamp: str) -> str:
    sanitized_job_name = re.sub(r"[^A-Za-z0-9_-]+", "_", (job_name or "").strip()).strip("_")
    return f"{sanitized_job_name or 'student_bulk_import'}_{timestamp}"


def _build_remaining_rows_filename_stem(job_name: str | None, timestamp: str) -> str:
    sanitized_job_name = re.sub(r"[^A-Za-z0-9_-]+", "_", (job_name or "").strip()).strip("_")
    return f"{sanitized_job_name or 'student_bulk_import'}_remaining_rows_{timestamp}"


def _build_glific_contact_file_name(
    tab_name: str,
    student_type: str,
    timestamp: str,
) -> str:
    sanitized_tab_name = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        (tab_name or "").strip(),
    ).strip("_")
    return f"{sanitized_tab_name or 'sheet'}_{student_type}_students_{timestamp}.csv"


def _glific_contact_row_for_headers(row: dict, headers: Iterable[str]) -> dict[str, str]:
    return {
        header: "" if row.get(header) is None else str(row.get(header))
        for header in headers
    }


def _create_and_upload_glific_contact_csvs(tab_names: list[str]) -> list[dict]:
    rows = _rows("""
        SELECT
            x.source_tab,
            x.source_priority,
            x.row_in_tab,
            COALESCE(x.existed_before_import, false) AS existed_before_import,
            COALESCE(s.name1, '') AS name,
            COALESCE(s.phone, '') AS phone,
            COALESCE(lang.language_name, '') AS language,
            '0' AS delete,
            COALESCE(sch.name, '') AS school_id,
            COALESCE(st.state_name, '') AS state,
            COALESCE(tm.mname, '') AS model,
            COALESCE(s.name1, '') AS buddy_name,
            COALESCE(b.batch_id, '') AS batch_id,
            COALESCE(NULLIF(x.grade_in, ''), NULLIF(s.grade, ''), '') AS grade,
            COALESCE(NULLIF(x.derived_level, ''), '') AS level,
            COALESCE(cv.name2, '') AS course
        FROM tmp_student_import_success x
        JOIN "tabStudent" s
          ON s.name = x.student_id
        LEFT JOIN "tabBatch" b
          ON b.name = x.batch_name
        LEFT JOIN "tabTAP Language" lang
          ON lang.name = s.language
        LEFT JOIN "tabSchool" sch
          ON sch.name = COALESCE(NULLIF(x.school_id, ''), NULLIF(s.school_id, ''))
        LEFT JOIN "tabState" st
          ON st.name = sch.state
        LEFT JOIN "tabTap Models" tm
          ON tm.name = sch.model
        LEFT JOIN "tabCourse Verticals" cv
          ON cv.name = x.vertical_id
        ORDER BY x.source_priority, x.row_in_tab
    """)

    rows_by_tab: dict[str, dict[str, list[dict]]] = {
        tab_name: {"new": [], "existing": []}
        for tab_name in tab_names
    }
    for row in rows:
        student_type = "existing" if row.get("existed_before_import") else "new"
        headers = (
            GLIFIC_EXISTING_STUDENT_CSV_HEADERS
            if student_type == "existing"
            else GLIFIC_NEW_STUDENT_CSV_HEADERS
        )
        contact_row = _glific_contact_row_for_headers(row, headers)
        rows_by_tab.setdefault(
            row["source_tab"],
            {"new": [], "existing": []},
        )[student_type].append(contact_row)

    timestamp = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
    exported_files: list[dict] = []
    for tab_name in tab_names:
        rows_for_tab = rows_by_tab.get(tab_name, {"new": [], "existing": []})
        for student_type, headers in (
            ("new", GLIFIC_NEW_STUDENT_CSV_HEADERS),
            ("existing", GLIFIC_EXISTING_STUDENT_CSV_HEADERS),
        ):
            contact_rows = rows_for_tab.get(student_type, [])
            file_name = _build_glific_contact_file_name(tab_name, student_type, timestamp)
            file_path = _upload_glific_contact_csv(
                file_name,
                contact_rows,
                headers=headers,
            )
            exported_files.append({
                "tab_name": tab_name,
                "student_type": student_type,
                "file_name": file_name,
                "file_path": file_path,
                "row_count": len(contact_rows),
            })
    return exported_files


def _stage_selected_tabs(
    workbook,
    tab_names: list[str],
    log_fn: Callable[[str], None] | None = None,
    chunk_size: int = 2000,
) -> int:
    missing_tabs = [tab_name for tab_name in tab_names if tab_name not in workbook.sheetnames]
    if missing_tabs:
        raise ValueError(
            f"Tabs not found in workbook: {missing_tabs}. "
            f"Available tabs: {workbook.sheetnames}"
        )

    total_rows = 0
    for priority, tab_name in enumerate(tab_names, start=1):
        ws = workbook[tab_name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            raise ValueError(f"Tab '{tab_name}' is empty")

        header_map = _resolve_header_map(tab_name, header)
        tab_rows: list[RawRow] = []
        tab_count = 0
        for row_in_tab, row_values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row_dict = {
                canonical_key: _stringify(row_values[idx] if idx < len(row_values) else None)
                for canonical_key, idx in header_map.items()
            }
            if not any(value.strip() for value in row_dict.values()):
                continue
            tab_rows.append(
                RawRow(
                    source_tab=tab_name,
                    source_priority=priority,
                    row_in_tab=row_in_tab,
                    student_name_raw=row_dict["student_name_raw"],
                    contact_no_raw=row_dict["contact_no_raw"],
                    gender_raw=row_dict["gender_raw"],
                    school_id_raw=row_dict["school_id_raw"],
                    language_raw=row_dict["language_raw"],
                    batch_raw=row_dict["batch_raw"],
                    grade_raw=row_dict["grade_raw"],
                    course_raw=row_dict["course_raw"],
                )
            )
            if len(tab_rows) >= chunk_size:
                _bulk_insert_stage(tab_rows, chunk_size=chunk_size)
                tab_count += len(tab_rows)
                total_rows += len(tab_rows)
                tab_rows = []
        if tab_rows:
            _bulk_insert_stage(tab_rows, chunk_size=chunk_size)
            tab_count += len(tab_rows)
            total_rows += len(tab_rows)
        _emit(log_fn, f"[student-import] staged tab='{tab_name}' raw_rows={tab_count}")
    return total_rows


def _resolve_header_map(tab_name: str, header: Iterable[object]) -> dict[str, int]:
    normalized = {str(value).strip(): idx for idx, value in enumerate(header) if value is not None}
    missing = [column for column in EXPECTED_COLUMNS if column not in normalized]
    if missing:
        raise ValueError(
            f"Tab '{tab_name}' is missing required columns: {missing}. "
            f"Found columns: {list(normalized.keys())}"
        )
    return {target_key: normalized[column_name] for column_name, target_key in EXPECTED_COLUMNS.items()}


def _stringify(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _prepare_temp_tables() -> None:
    sql = """
    DROP TABLE IF EXISTS tmp_student_import_raw;
    DROP TABLE IF EXISTS tmp_student_import_clean;
    DROP TABLE IF EXISTS tmp_student_import_effective;
    DROP TABLE IF EXISTS tmp_student_import_validation;
    DROP TABLE IF EXISTS tmp_student_import_valid;
    DROP TABLE IF EXISTS tmp_student_import_invalid;
    DROP TABLE IF EXISTS tmp_student_import_match;
    DROP TABLE IF EXISTS tmp_student_import_to_insert;
    DROP TABLE IF EXISTS tmp_student_import_inserted;
    DROP TABLE IF EXISTS tmp_student_import_resolved;
    DROP TABLE IF EXISTS tmp_student_import_success;

    CREATE TEMP TABLE tmp_student_import_raw (
        source_tab text NOT NULL,
        source_priority integer NOT NULL,
        row_in_tab integer NOT NULL,
        student_name_raw text,
        contact_no_raw text,
        gender_raw text,
        school_id_raw text,
        language_raw text,
        batch_raw text,
        grade_raw text,
        course_raw text
    );

    CREATE TEMP TABLE tmp_student_import_success (
        source_tab text NOT NULL,
        source_priority integer NOT NULL,
        row_in_tab integer NOT NULL,
        student_id text NOT NULL,
        school_id text,
        grade_in text,
        derived_level text,
        batch_name text,
        vertical_id text,
        existed_before_import boolean NOT NULL DEFAULT false
    );
    """
    frappe.db.sql(sql)


def _bulk_insert_stage(rows: list[RawRow], chunk_size: int = 2000) -> None:
    conn = getattr(frappe.db, "_conn", None)
    if conn is None:
        raise RuntimeError("Frappe database connection is not available")

    insert_sql = """
        INSERT INTO tmp_student_import_raw (
            source_tab,
            source_priority,
            row_in_tab,
            student_name_raw,
            contact_no_raw,
            gender_raw,
            school_id_raw,
            language_raw,
            batch_raw,
            grade_raw,
            course_raw
        ) VALUES %s
    """
    values = [
        (
            row.source_tab,
            row.source_priority,
            row.row_in_tab,
            row.student_name_raw,
            row.contact_no_raw,
            row.gender_raw,
            row.school_id_raw,
            row.language_raw,
            row.batch_raw,
            row.grade_raw,
            row.course_raw,
        )
        for row in rows
    ]

    with conn.cursor() as cursor:
        for start in range(0, len(values), chunk_size):
            execute_values(cursor, insert_sql, values[start:start + chunk_size], page_size=chunk_size)


def _build_clean_stage(sample_test: int) -> None:
    sql = f"""
    CREATE TEMP TABLE tmp_student_import_clean AS
    WITH base AS (
        SELECT
            source_tab,
            source_priority,
            row_in_tab,
            trim(coalesce(student_name_raw, '')) AS student_name_raw,
            trim(coalesce(contact_no_raw, '')) AS contact_no_raw,
            trim(coalesce(school_id_raw, '')) AS school_id_raw,
            trim(coalesce(language_raw, '')) AS language_raw,
            trim(coalesce(batch_raw, '')) AS batch_raw,
            trim(coalesce(grade_raw, '')) AS grade_raw,
            trim(coalesce(course_raw, '')) AS course_raw,
            regexp_replace(
                regexp_replace(trim(coalesce(contact_no_raw, '')), '\\.0+$', ''),
                '\\D',
                '',
                'g'
            ) AS phone_digits,
            trim(coalesce(gender_raw, '')) AS gender_raw,
            nullif(trim(coalesce(school_id_raw, '')), '') AS school_id_in,
            nullif(trim(coalesce(language_raw, '')), '') AS language_name_in,
            nullif(trim(coalesce(batch_raw, '')), '') AS batch_in,
            nullif(trim(coalesce(grade_raw, '')), '') AS grade_in,
            nullif(trim(coalesce(course_raw, '')), '') AS course_name_in
        FROM tmp_student_import_raw
    ),
    phones AS (
        SELECT
            *,
            CASE
                WHEN length(phone_digits) = 10 THEN '91' || phone_digits
                WHEN length(phone_digits) = 12 AND left(phone_digits, 2) = '91' THEN phone_digits
                ELSE NULL
            END AS phone_12,
            CASE
                WHEN length(phone_digits) = 10 THEN phone_digits
                WHEN length(phone_digits) = 12 AND left(phone_digits, 2) = '91' THEN right(phone_digits, 10)
                ELSE NULL
            END AS phone_10
        FROM base
    ),
    cleaned AS (
        SELECT
            *,
            btrim(
                regexp_replace(
                    regexp_replace(student_name_raw, '[^A-Za-z ]+', '', 'g'),
                    '\\s+',
                    ' ',
                    'g'
                )
            ) AS cleaned_name_tmp,
            CASE upper(gender_raw)
                WHEN 'M' THEN 'Male'
                WHEN 'MALE' THEN 'Male'
                WHEN 'F' THEN 'Female'
                WHEN 'FEMALE' THEN 'Female'
                ELSE NULL
            END AS normalized_gender
        FROM phones
    )
    SELECT
        source_tab,
        source_priority,
        row_in_tab,
        student_name_raw,
        contact_no_raw,
        school_id_raw,
        language_raw,
        batch_raw,
        grade_raw,
        course_raw,
        CASE
            WHEN cleaned_name_tmp IS NULL OR cleaned_name_tmp = '' THEN 'Champ'
            ELSE cleaned_name_tmp
        END AS cleaned_name,
        phone_12,
        phone_10,
        gender_raw,
        normalized_gender,
        school_id_in,
        language_name_in,
        batch_in,
        CASE
            WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
            ELSE grade_in
        END AS grade_in,
        CASE
            WHEN lower(coalesce(course_name_in, '')) = 'science' THEN 'Science Lab'
            ELSE course_name_in
        END AS course_name_in,
        CASE
            WHEN (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            ) ~ '^\\d+$'
            AND (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            )::int BETWEEN 1 AND 3 THEN 'Level 0'
            WHEN (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            ) ~ '^\\d+$'
            AND (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            )::int BETWEEN 4 AND 5 THEN 'Level 1'
            WHEN (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            ) ~ '^\\d+$'
            AND (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            )::int BETWEEN 6 AND 8 THEN 'Level 2'
            WHEN (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            ) ~ '^\\d+$'
            AND (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            )::int BETWEEN 9 AND 10 THEN 'Level 3'
            WHEN (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            ) ~ '^\\d+$'
            AND (
                CASE
                    WHEN grade_in ~ '^\\d+\\.0+$' THEN split_part(grade_in, '.', 1)
                    ELSE grade_in
                END
            )::int BETWEEN 11 AND 12 THEN 'Level 4'
            ELSE NULL
        END AS derived_level
    FROM cleaned
    """
    frappe.db.sql(sql)
    
    
def _build_validation_tables(sample_test: int) -> None:
    sample_limit_sql = ""
    params: tuple[object, ...] = ()
    if sample_test > 0:
        sample_limit_sql = "LIMIT %s"
        params = (sample_test,)

    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_validation")
    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_valid")
    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_invalid")

    sql = f"""
    CREATE TEMP TABLE tmp_student_import_validation AS
    WITH sampled AS (
        SELECT *
        FROM tmp_student_import_clean
        ORDER BY source_priority, row_in_tab
        {sample_limit_sql}
    ),
    ranked AS (
        SELECT
            s.*,
            CASE
                WHEN s.phone_12 IS NOT NULL THEN
                    row_number() OVER (PARTITION BY s.phone_12 ORDER BY s.source_priority, s.row_in_tab)
                ELSE 1
            END AS phone_rank,
            first_value(s.cleaned_name) OVER (
                PARTITION BY s.phone_12
                ORDER BY s.source_priority, s.row_in_tab
            ) AS kept_cleaned_name,
            first_value(s.phone_12) OVER (
                PARTITION BY s.phone_12
                ORDER BY s.source_priority, s.row_in_tab
            ) AS kept_phone_12,
            first_value(s.source_tab) OVER (
                PARTITION BY s.phone_12
                ORDER BY s.source_priority, s.row_in_tab
            ) AS kept_source_tab,
            first_value(s.row_in_tab) OVER (
                PARTITION BY s.phone_12
                ORDER BY s.source_priority, s.row_in_tab
            ) AS kept_row_in_tab
        FROM sampled s
    ),
    enriched AS (
        SELECT
            r.*,
            EXISTS (
                SELECT 1
                FROM "tabSchool" s
                WHERE s.name = r.school_id_in
                   OR substring(s.name from '(SC[0-9]+)$') = r.school_id_in
            ) AS school_exists,
            EXISTS (
                SELECT 1
                FROM "tabBatch" b
                WHERE b.name = r.batch_in
                   OR b.batch_id = r.batch_in
            ) AS batch_exists,
            EXISTS (
                SELECT 1
                FROM "tabTAP Language" l
                WHERE l.language_name = r.language_name_in
            ) AS language_exists,
            EXISTS (
                SELECT 1
                FROM "tabCourse Verticals" cv
                WHERE cv.name2 = r.course_name_in
                   OR cv.name = r.course_name_in
                   OR cv.vertical_id = r.course_name_in
            ) AS vertical_exists,
            COALESCE((
                SELECT count(*)
                FROM "tabStudent" st
                WHERE st.phone IN (r.phone_12, r.phone_10)
            ), 0) AS existing_phone_match_count
        FROM ranked r
    )
    SELECT
        *,
        (phone_12 IS NULL) AS fail_phone,
        (phone_12 IS NOT NULL AND phone_rank > 1) AS fail_duplicate_phone,
        (derived_level IS NULL) AS fail_grade,
        (NOT school_exists) AS fail_school,
        (NOT batch_exists) AS fail_batch,
        (NOT language_exists) AS fail_language,
        (NOT vertical_exists) AS fail_course,
        false AS fail_ambiguous_phone,
        concat_ws(
            '; ',
            CASE WHEN phone_12 IS NULL THEN 'Invalid Contact No.' END,
            CASE
                WHEN phone_12 IS NOT NULL AND phone_rank > 1 THEN
                    'Duplicate Contact No. in selected import set. '
                    || 'Kept student: '
                    || coalesce(kept_cleaned_name, 'Champ')
                    || ' ('
                    || coalesce(kept_phone_12, '')
                    || ') from '
                    || coalesce(kept_source_tab, '')
                    || ' row '
                    || coalesce(kept_row_in_tab::text, '')
            END,
            CASE WHEN derived_level IS NULL THEN 'Invalid Grade' END,
            CASE WHEN NOT school_exists THEN 'School ID not found in this site' END,
            CASE WHEN NOT batch_exists THEN 'Batch not found in this site' END,
            CASE WHEN NOT language_exists THEN 'Language not found in this site' END,
            CASE WHEN NOT vertical_exists THEN 'Course not found in this site' END
        ) AS validation_errors,
        NOT (
            (phone_12 IS NULL)
            OR (phone_12 IS NOT NULL AND phone_rank > 1)
            OR (derived_level IS NULL)
            OR (NOT school_exists)
            OR (NOT batch_exists)
            OR (NOT language_exists)
            OR (NOT vertical_exists)
        ) AS is_valid
    FROM enriched
    """
    frappe.db.sql(sql, params)

    frappe.db.sql("""
        CREATE TEMP TABLE tmp_student_import_valid AS
        SELECT *
        FROM tmp_student_import_validation
        WHERE is_valid
        ORDER BY source_priority, row_in_tab
    """)

    frappe.db.sql("""
        CREATE TEMP TABLE tmp_student_import_invalid AS
        SELECT *
        FROM tmp_student_import_validation
        WHERE NOT is_valid
        ORDER BY source_priority, row_in_tab
    """)


def _run_prechecks() -> dict:
    return {
        "raw_rows": _scalar("SELECT count(*) FROM tmp_student_import_raw"),
        "effective_rows": _scalar("SELECT count(*) FROM tmp_student_import_valid"),
        "valid_rows": _scalar("SELECT count(*) FROM tmp_student_import_valid"),
        "invalid_rows": _scalar("SELECT count(*) FROM tmp_student_import_invalid"),
        "invalid_phone_rows": _scalar("SELECT count(*) FROM tmp_student_import_invalid WHERE fail_phone"),
        "duplicate_input_phones": _scalar("""
            SELECT count(*)
            FROM tmp_student_import_invalid
            WHERE fail_duplicate_phone
        """),
        "ignored_gender_rows": _scalar("""
            SELECT count(*)
            FROM tmp_student_import_validation
            WHERE normalized_gender IS NULL
        """),
        "invalid_grade_rows": _scalar("SELECT count(*) FROM tmp_student_import_invalid WHERE fail_grade"),
        "missing_school_rows": _scalar("SELECT count(*) FROM tmp_student_import_invalid WHERE fail_school"),
        "missing_batch_rows": _scalar("SELECT count(*) FROM tmp_student_import_invalid WHERE fail_batch"),
        "missing_language_rows": _scalar("SELECT count(*) FROM tmp_student_import_invalid WHERE fail_language"),
        "missing_vertical_rows": _scalar("SELECT count(*) FROM tmp_student_import_invalid WHERE fail_course"),
        "ambiguous_existing_phone_matches": _scalar("""
            SELECT count(*)
            FROM tmp_student_import_validation
            WHERE existing_phone_match_count > 1
        """),
        "missing_school_examples": _rows("""
            SELECT school_id_in AS school_id, count(*) AS row_count
            FROM tmp_student_import_invalid
            WHERE fail_school
            GROUP BY school_id_in
            ORDER BY count(*) DESC, school_id_in ASC
            LIMIT 10
        """),
        "missing_batch_examples": _rows("""
            SELECT batch_in AS batch, count(*) AS row_count
            FROM tmp_student_import_invalid
            WHERE fail_batch
            GROUP BY batch_in
            ORDER BY count(*) DESC, batch_in ASC
            LIMIT 10
        """),
    }


def _print_precheck(precheck: dict, log_fn: Callable[[str], None] | None = None) -> None:
    _emit(log_fn, "[student-import] precheck")
    for key, value in precheck.items():
        if key.endswith("_examples"):
            if value:
                _emit(log_fn, f"  - {key}:")
                for row in value:
                    _emit(log_fn, f"      {row}")
            continue
        _emit(log_fn, f"  - {key}: {value}")


def _raise_if_blocking_precheck(precheck: dict) -> None:
    return None


def _prepare_batch_subset(offset: int, batch_size: int) -> None:
    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_batch")
    frappe.db.sql("""
        CREATE TEMP TABLE tmp_student_import_batch AS
        SELECT *
        FROM tmp_student_import_valid
        ORDER BY source_priority, row_in_tab
        OFFSET %s
        LIMIT %s
    """, (offset, batch_size))


def _execute_import(import_date: date, import_user: str, source_table: str = "tmp_student_import_valid") -> dict:
    source_sql = f"""
        CREATE TEMP TABLE tmp_student_import_match AS
        SELECT
            e.source_tab,
            e.source_priority,
            e.row_in_tab,
            e.cleaned_name,
            e.phone_12,
            e.phone_10,
            e.normalized_gender,
            school.name AS school_id,
            e.grade_in,
            batch.name AS batch_name,
            e.derived_level,
            l.name AS language_id,
            cv.name AS vertical_id,
            s.name AS student_id
        FROM {source_table} e
        JOIN "tabTAP Language" l
          ON l.language_name = e.language_name_in
        JOIN LATERAL (
            SELECT cv.name, cv.name2, cv.vertical_id
            FROM "tabCourse Verticals" cv
            WHERE cv.name2 = e.course_name_in
               OR cv.name = e.course_name_in
               OR cv.vertical_id = e.course_name_in
            ORDER BY cv.name
            LIMIT 1
        ) cv ON true
        JOIN LATERAL (
            SELECT s2.name
            FROM "tabSchool" s2
            WHERE s2.name = e.school_id_in
               OR substring(s2.name from '(SC[0-9]+)$') = e.school_id_in
            ORDER BY s2.name
            LIMIT 1
        ) school ON true
        JOIN LATERAL (
            SELECT b.name
            FROM "tabBatch" b
            WHERE b.name = e.batch_in
               OR b.batch_id = e.batch_in
            ORDER BY b.name
            LIMIT 1
        ) batch ON true
        LEFT JOIN LATERAL (
            SELECT s.name
            FROM "tabStudent" s
            LEFT JOIN LATERAL (
                SELECT max(se.date_joining) AS latest_date_joining
                FROM "tabStudent Enrollment" se
                WHERE se.parent = s.name
            ) latest_enrollment ON true
            WHERE s.phone IN (e.phone_12, e.phone_10)
            ORDER BY
                CASE
                    WHEN lower(
                        CASE
                            WHEN btrim(
                                regexp_replace(
                                    regexp_replace(coalesce(s.name1, ''), '[^A-Za-z ]+', '', 'g'),
                                    '\\s+',
                                    ' ',
                                    'g'
                                )
                            ) = '' THEN 'Champ'
                            ELSE btrim(
                                regexp_replace(
                                    regexp_replace(coalesce(s.name1, ''), '[^A-Za-z ]+', '', 'g'),
                                    '\\s+',
                                    ' ',
                                    'g'
                                )
                            )
                        END
                    ) = lower(e.cleaned_name)
                    THEN 1 ELSE 0
                END DESC,
                latest_enrollment.latest_date_joining DESC NULLS LAST,
                lower(coalesce(s.name1, '')) ASC,
                s.name ASC
            LIMIT 1
        ) s ON true
    """
    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_match")
    frappe.db.sql(source_sql)

    updated_count = _update_existing_students(import_user=import_user)
    inserted_count = _insert_new_students(import_date=import_date, import_user=import_user)
    enrollment_summary = _insert_enrollments(import_date=import_date, import_user=import_user)

    return {
        "batch_rows": _scalar(f"SELECT count(*) FROM {source_table}"),
        "updated_students": updated_count,
        "inserted_students": inserted_count,
        "inserted_enrollments": enrollment_summary["inserted_enrollments"],
        "skipped_duplicate_enrollments": enrollment_summary["skipped_duplicate_enrollments"],
    }


def _update_existing_students(import_user: str) -> int:
    rows = frappe.db.sql("""
        UPDATE "tabStudent" st
           SET name1 = m.cleaned_name,
               phone = m.phone_12,
               school_id = m.school_id,
               language = m.language_id,
               grade = m.grade_in,
               whatsapp_consent = 1,
               gender = CASE
                           WHEN m.normalized_gender IS NOT NULL THEN m.normalized_gender
                           ELSE st.gender
                        END,
               modified = now(),
               modified_by = %s
          FROM tmp_student_import_match m
         WHERE st.name = m.student_id
    """, (import_user,))
    return rows if isinstance(rows, int) else _scalar("""
        SELECT count(*)
        FROM tmp_student_import_match
        WHERE student_id IS NOT NULL
    """)


def _insert_new_students(import_date: date, import_user: str) -> int:
    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_to_insert")
    frappe.db.sql("""
        CREATE TEMP TABLE tmp_student_import_to_insert AS
        SELECT *
        FROM tmp_student_import_match
        WHERE student_id IS NULL
    """)

    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_inserted")
    total_to_insert = _scalar("SELECT count(*) FROM tmp_student_import_to_insert")
    if not total_to_insert:
        frappe.db.sql("""
            CREATE TEMP TABLE tmp_student_import_inserted (
                student_id text,
                phone_12 text
            )
        """)
        return 0

    _sync_student_series_counter()

    frappe.db.sql("""
        CREATE TEMP TABLE tmp_student_import_inserted (
            student_id text,
            phone_12 text
        )
    """)

    frappe.db.sql("""
        WITH cnt AS (
            SELECT count(*) AS n
            FROM tmp_student_import_to_insert
        ),
        bump AS (
            UPDATE "tabSeries"
               SET current = current + (SELECT n FROM cnt)
             WHERE name = 'ST'
         RETURNING current
        ),
        numbered AS (
            SELECT
                t.*,
                row_number() OVER (ORDER BY source_priority, row_in_tab) AS rn,
                (SELECT current FROM bump) - (SELECT n FROM cnt) AS base_no
            FROM tmp_student_import_to_insert t
        ),
        inserted AS (
            INSERT INTO "tabStudent" (
                name,
                creation,
                modified,
                modified_by,
                owner,
                docstatus,
                name1,
                phone,
                gender,
                school_id,
                grade,
                joined_on,
                status,
                language,
                whatsapp_consent
            )
            SELECT
                'ST' || lpad((base_no + rn)::text, 8, '0') AS name,
                now(),
                now(),
                %(import_user)s,
                %(import_user)s,
                0,
                cleaned_name,
                phone_12,
                normalized_gender,
                school_id,
                grade_in,
                %(import_date)s,
                'active',
                language_id,
                0
            FROM numbered
            RETURNING name, phone
        )
        INSERT INTO tmp_student_import_inserted (student_id, phone_12)
        SELECT name, phone
        FROM inserted
    """, {
        "import_user": import_user,
        "import_date": import_date,
    })

    return total_to_insert


def _sync_student_series_counter() -> None:
    """Self-heal Student series counters from real Student rows.

    Legacy meta `format:ST{########}` uses the empty-string series key,
    while corrected meta `format:ST.########` uses the `ST` series key.
    Keep both aligned so non-bulk Student inserts remain safe after import.
    """
    frappe.db.sql("""
        INSERT INTO "tabSeries" (name, current)
        VALUES ('ST', 0), ('', 0)
        ON CONFLICT (name) DO NOTHING
    """)
    frappe.db.sql("""
        WITH max_student AS (
            SELECT COALESCE(
                max(CASE
                    WHEN name ~ '^ST[0-9]{8}$' THEN substring(name from 3)::integer
                    ELSE 0
                END),
                0
            ) AS max_no
            FROM "tabStudent"
        )
        UPDATE "tabSeries" ts
           SET current = GREATEST(ts.current, ms.max_no)
          FROM max_student ms
         WHERE ts.name IN ('ST', '')
    """)


def _insert_enrollments(import_date: date, import_user: str) -> dict:
    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_resolved")
    frappe.db.sql("DROP TABLE IF EXISTS tmp_student_import_enrollment_to_insert")
    frappe.db.sql("""
        CREATE TEMP TABLE tmp_student_import_resolved AS
        SELECT
            m.source_tab,
            m.source_priority,
            m.row_in_tab,
            m.phone_12,
            m.school_id,
            m.grade_in,
            m.batch_name,
            m.derived_level,
            m.vertical_id,
            m.student_id,
            true AS existed_before_import
        FROM tmp_student_import_match m
        WHERE m.student_id IS NOT NULL

        UNION ALL

        SELECT
            t.source_tab,
            t.source_priority,
            t.row_in_tab,
            t.phone_12,
            t.school_id,
            t.grade_in,
            t.batch_name,
            t.derived_level,
            t.vertical_id,
            i.student_id,
            false AS existed_before_import
        FROM tmp_student_import_to_insert t
        JOIN tmp_student_import_inserted i
          ON i.phone_12 = t.phone_12
    """)

    total_rows = _scalar("SELECT count(*) FROM tmp_student_import_resolved")
    if not total_rows:
        return {
            "inserted_enrollments": 0,
            "skipped_duplicate_enrollments": 0,
        }

    frappe.db.sql("""
        CREATE TEMP TABLE tmp_student_import_enrollment_to_insert AS
        SELECT r.*
        FROM tmp_student_import_resolved r
        WHERE NOT EXISTS (
            SELECT 1
            FROM "tabStudent Enrollment" existing
            WHERE existing.parent = r.student_id
              AND existing.parenttype = 'Student'
              AND existing.parentfield = 'enrollment'
              AND existing.batch IS NOT DISTINCT FROM r.batch_name
              AND existing.vertical IS NOT DISTINCT FROM r.vertical_id
              AND existing.level IS NOT DISTINCT FROM r.derived_level
              AND existing.grade IS NOT DISTINCT FROM r.grade_in
              AND existing.school IS NOT DISTINCT FROM r.school_id
        )
    """)

    enrollment_rows_to_insert = _scalar("SELECT count(*) FROM tmp_student_import_enrollment_to_insert")

    frappe.db.sql("""
        INSERT INTO "tabStudent Enrollment" (
            name,
            creation,
            modified,
            modified_by,
            owner,
            docstatus,
            parent,
            parentfield,
            parenttype,
            idx,
            batch,
            vertical,
            level,
            grade,
            date_joining,
            school,
            whatsapp_response
        )
        SELECT
            md5(
                clock_timestamp()::text
                || random()::text
                || r.student_id
                || coalesce(r.batch_name, '')
                || r.source_tab
                || r.row_in_tab::text
            ) AS name,
            now(),
            now(),
            %(import_user)s,
            %(import_user)s,
            0,
            r.student_id,
            'enrollment',
            'Student',
            coalesce((
                SELECT max(e.idx)
                FROM "tabStudent Enrollment" e
                WHERE e.parent = r.student_id
            ), 0) + 1 AS idx,
            r.batch_name,
            r.vertical_id,
            r.derived_level,
            r.grade_in,
            %(import_date)s,
            r.school_id,
            0
        FROM tmp_student_import_enrollment_to_insert r
    """, {
        "import_user": import_user,
        "import_date": import_date,
    })
    frappe.db.sql("""
        INSERT INTO tmp_student_import_success (
            source_tab,
            source_priority,
            row_in_tab,
            student_id,
            school_id,
            grade_in,
            derived_level,
            batch_name,
            vertical_id,
            existed_before_import
        )
        SELECT
            source_tab,
            source_priority,
            row_in_tab,
            student_id,
            school_id,
            grade_in,
            derived_level,
            batch_name,
            vertical_id,
            existed_before_import
        FROM tmp_student_import_resolved
    """)
    return {
        "inserted_enrollments": enrollment_rows_to_insert,
        "skipped_duplicate_enrollments": total_rows - enrollment_rows_to_insert,
    }


def _scalar(sql: str, params: object | None = None) -> int:
    result = frappe.db.sql(sql, params)
    return int(result[0][0]) if result else 0


def _rows(sql: str, params: object | None = None) -> list[dict]:
    return frappe.db.sql(sql, params, as_dict=True) or []
